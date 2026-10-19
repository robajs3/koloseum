"""Podgląd dokumentów (services/preview_service.py + endpointy preview-doc)."""
import base64
import io
import tempfile
import unittest
from unittest import mock

from app import create_app
from config import Config
from models import db, User, Room, RoomMember, Subject, StudyMaterial
from services import preview_service as ps
from tests import fixtures as fx
from tests.mock_filevault import MockFileVault

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
W = fx.W
R = fx.R


def render(data: bytes, name: str) -> str:
    with tempfile.NamedTemporaryFile(suffix=name[name.rfind("."):]) as f:
        f.write(data); f.flush()
        kind, html = ps.render_document(f.name, name)
    assert kind == "html"
    return html


def rich_docx() -> bytes:
    body = (
        f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Tytuł rozdziału</w:t></w:r></w:p>'
        f'<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>pogrubione</w:t></w:r><w:r><w:t xml:space="preserve"> zwykłe </w:t></w:r>'
        f'<w:r><w:t>&lt;script&gt;alert(1)&lt;/script&gt;</w:t></w:r></w:p>'
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr><w:r><w:t>punkt A</w:t></w:r></w:p>'
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr></w:pPr><w:r><w:t>krok 1</w:t></w:r></w:p>'
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr></w:pPr><w:r><w:t>krok 2</w:t></w:r></w:p>'
        f'<w:p><w:hyperlink r:id="rId5"><w:r><w:t>dobry link</w:t></w:r></w:hyperlink>'
        f'<w:hyperlink r:id="rId6"><w:r><w:t>zły link</w:t></w:r></w:hyperlink></w:p>'
        f'<w:p><w:r><w:drawing><a:blip xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" r:embed="rId7"/></w:drawing></w:r></w:p>'
        f'<w:tbl><w:tr><w:tc><w:p><w:r><w:t>komórka1</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>komórka2</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
    )
    return fx._zip({
        "word/document.xml": f'<?xml version="1.0"?><w:document xmlns:w="{W}" xmlns:r="{R}"><w:body>{body}</w:body></w:document>',
        "word/styles.xml": f'<w:styles xmlns:w="{W}"><w:style w:styleId="Heading1"><w:name w:val="heading 1"/></w:style></w:styles>',
        "word/numbering.xml": (
            f'<w:numbering xmlns:w="{W}">'
            '<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:numFmt w:val="bullet"/></w:lvl></w:abstractNum>'
            '<w:abstractNum w:abstractNumId="1"><w:lvl w:ilvl="0"><w:numFmt w:val="decimal"/></w:lvl></w:abstractNum>'
            '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num><w:num w:numId="2"><w:abstractNumId w:val="1"/></w:num>'
            '</w:numbering>'),
        "word/_rels/document.xml.rels": (
            f'<Relationships xmlns="{fx.PR}">'
            '<Relationship Id="rId5" Target="https://example.com/ok" TargetMode="External"/>'
            '<Relationship Id="rId6" Target="javascript:alert(1)" TargetMode="External"/>'
            '<Relationship Id="rId7" Target="media/p.png"/></Relationships>'),
        "word/media/p.png": PNG,
    })


class RenderTests(unittest.TestCase):
    def test_kinds(self):
        k = ps.preview_kind
        for n, expected in [("a.PNG", "image"), ("a.avif", "image"), ("a.pdf", "pdf"), ("a.mov", "video"),
                            ("a.flac", "audio"), ("a.py", "text"), ("a.toml", "text"), ("a.docx", "office"),
                            ("a.pptx", "office"), ("a.xlsx", "office"), ("a.odt", "office"), ("a.csv", "office"),
                            ("a.ipynb", "office"), ("a.rtf", "office"), ("a.zip", None), ("noext", None)]:
            self.assertEqual(k(n), expected, n)
        with mock.patch("services.preview_service.has_libreoffice", return_value=False):
            self.assertIsNone(k("old.doc"))
        with mock.patch("services.preview_service.has_libreoffice", return_value=True):
            self.assertEqual(k("old.xls"), "office")

    def test_docx_rich(self):
        h = render(rich_docx(), "a.docx")
        self.assertIn("<h1>Tytuł rozdziału</h1>", h)
        self.assertIn("<b>pogrubione</b>", h)
        self.assertNotIn("<script>alert", h)                 # treść escapowana
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", h)
        self.assertIn("•", h)                                # lista punktowana
        self.assertIn("1.", h); self.assertIn("2.", h)       # lista numerowana
        self.assertIn('href="https://example.com/ok"', h)
        self.assertNotIn("javascript:", h)                   # niebezpieczny link odrzucony
        self.assertIn("zły link", h)                         # ale tekst zostaje
        self.assertIn("data:image/png;base64,", h)
        self.assertIn("<td>", h); self.assertIn("komórka2", h)

    def test_simple_fixture_docx_pptx_xlsx_odt(self):
        self.assertIn("Zażółć gęślą jaźń", render(fx.docx(["Zażółć gęślą jaźń"]), "a.docx"))
        h = render(fx.pptx([("slide2", "Drugi"), ("slide1", "Pierwszy")]), "a.pptx")
        self.assertLess(h.index("Drugi"), h.index("Pierwszy"))   # kolejność z presentation.xml
        self.assertIn("Slajd 2", h)
        h = render(fx.xlsx({"Oceny": [["Imię", "Punkty"], ["Ala", 12.5], ["Ola", True]], "Druga": [["x"]]}), "a.xlsx")
        self.assertIn("Oceny", h); self.assertIn("Druga", h)
        self.assertIn("<td>12.5</td>", h); self.assertIn("TRUE", h)
        self.assertIn("Dokument ODF", render(fx.odt(["Dokument ODF"]), "a.odt"))

    def test_xlsx_dates_and_percent(self):
        styles = (f'<styleSheet xmlns="{fx.S}"><cellXfs><xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="9"/></cellXfs></styleSheet>')
        sheet = (f'<worksheet xmlns="{fx.S}"><sheetData><row r="1">'
                 '<c r="A1" s="1"><v>45000</v></c><c r="B1" s="2"><v>0.256</v></c><c r="C1"><v>0.30000000000000004</v></c>'
                 '</row></sheetData></worksheet>')
        data = fx._zip({
            "xl/workbook.xml": f'<workbook xmlns="{fx.S}" xmlns:r="{fx.R}"><sheets><sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>',
            "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{fx.PR}"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>',
            "xl/worksheets/sheet1.xml": sheet, "xl/styles.xml": styles})
        h = render(data, "d.xlsx")
        self.assertIn("2023-03-15", h); self.assertIn("25.6%", h); self.assertIn("<td>0.3</td>", h)

    def test_csv_semicolon_cp1250_and_ipynb_and_rtf(self):
        h = render("imię;wiek\nŁukasz;20\n".encode("cp1250"), "a.csv")
        self.assertIn("<th>imię</th>", h); self.assertIn("Łukasz", h)
        nb = ('{"cells":[{"cell_type":"markdown","source":["# Tytuł"]},'
              '{"cell_type":"code","source":["print(1)"],"outputs":[{"output_type":"stream","text":["1\\n"]}]}]}')
        h = render(nb.encode(), "a.ipynb")
        self.assertIn("<h1>Tytuł</h1>", h); self.assertIn("print(1)", h)
        h = render(br"{\rtf1\ansi{\fonttbl{\f0 Arial;}}Hello \b world\b0\par Drugi \'a3ukasz}", "a.rtf")
        self.assertIn("Hello", h); self.assertIn("Łukasz", h); self.assertNotIn("Arial", h.split("<body>")[1])  # tabela fontów RTF nie wycieka

    def test_errors(self):
        for data, name, status in [(b"to nie zip", "a.docx", 422), (b"x", "a.unknown", 415)]:
            with tempfile.NamedTemporaryFile(suffix=name[name.rfind("."):]) as f:
                f.write(data); f.flush()
                with self.assertRaises(ps.PreviewError) as cm:
                    ps.render_document(f.name, name)
                self.assertEqual(cm.exception.status, status)

    def test_xml_entities_rejected(self):
        evil = fx._zip({"word/document.xml": '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><w:document xmlns:w="%s"/>' % W})
        with tempfile.NamedTemporaryFile(suffix=".docx") as f:
            f.write(evil); f.flush()
            with self.assertRaises(ps.PreviewError):
                ps.render_document(f.name, "a.docx")

    def test_output_limit_truncates(self):
        with mock.patch.object(ps, "MAX_HTML_BYTES", 3000):
            h = render(fx.docx(["akapit numer %d z dłuższą treścią" % i for i in range(500)]), "a.docx")
        self.assertIn("skrócony", h)
        self.assertIn("akapit numer 0", h)
        self.assertNotIn("akapit numer 499", h)

    def test_legacy_without_libreoffice(self):
        with tempfile.NamedTemporaryFile(suffix=".doc") as f:
            f.write(b"\xd0\xcf\x11\xe0" + b"0" * 50); f.flush()
            with mock.patch("shutil.which", return_value=None):
                with self.assertRaises(ps.PreviewError) as cm:
                    ps.render_document(f.name, "a.doc")
            self.assertEqual(cm.exception.status, 415)

    @unittest.skipUnless(__import__("importlib").util.find_spec("pptx") and __import__("importlib").util.find_spec("openpyxl"),
                         "python-pptx / openpyxl niedostępne")
    def test_real_office_generated_files(self):
        import openpyxl
        from pptx import Presentation
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Dane"
        ws.append(["Imię", "Punkty", "Data"]); ws.append(["Ala", 12.5, __import__("datetime").date(2024, 5, 17)])
        buf = io.BytesIO(); wb.save(buf)
        h = render(buf.getvalue(), "r.xlsx")
        self.assertIn("Ala", h); self.assertIn("2024-05-17", h)
        prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Tytuł slajdu"; s.placeholders[1].text = "Treść punktu"
        s.notes_slide.notes_text_frame.text = "Moja notatka"
        buf = io.BytesIO(); prs.save(buf)
        h = render(buf.getvalue(), "r.pptx")
        self.assertIn("Tytuł slajdu", h); self.assertIn("Treść punktu", h); self.assertIn("Moja notatka", h)


class EndpointTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.fv = MockFileVault().start()
        cls.fv.file("/filevault/s/DOCX", fx.docx(["Treść z FileVault"]), "plik.docx")
        cls.fv.file("/filevault/s/PNG", b"\x89PNG" + b"0" * 100, "foto.png")

        class Cfg(Config):
            TESTING = True
            SECRET_KEY = "test"
            SQLALCHEMY_DATABASE_URI = f"sqlite:///{cls.tmp.name}/t.db"
            UPLOAD_FOLDER = f"{cls.tmp.name}/uploads"
            PREFIX = ""
            FILEVAULT_INTERNAL_URL = cls.fv.url

        cls.app = create_app(Cfg)
        import os
        os.makedirs(f"{cls.tmp.name}/uploads/materials", exist_ok=True)
        with cls.app.app_context():
            db.create_all()
            a = User(username="a", email="a@x.pl"); a.set_password("x")
            o = User(username="o", email="o@x.pl"); o.set_password("x")
            db.session.add_all([a, o]); db.session.flush()
            r1 = Room(name="R1", created_by=a.id); r2 = Room(name="R2", created_by=o.id)
            db.session.add_all([r1, r2]); db.session.flush()
            db.session.add_all([RoomMember(room_id=r1.id, user_id=a.id, role="admin"),
                                RoomMember(room_id=r2.id, user_id=o.id, role="admin")])
            s1 = Subject(name="S1", room_id=r1.id, created_by=a.id)
            s2 = Subject(name="S2", room_id=r2.id, created_by=o.id)
            db.session.add_all([s1, s2]); db.session.flush()

            def mat(subj, title, fname, data):
                with open(f"{cls.tmp.name}/uploads/materials/{fname}", "wb") as f:
                    f.write(data)
                m = StudyMaterial(subject_id=subj.id, uploaded_by=a.id, title=title,
                                  filename=fname, original_filename=title)
                db.session.add(m); db.session.flush()
                return m.id
            cls.docx_id = mat(s1, "wyklad.docx", "u1.docx", fx.docx(["Treść wykładu"]))
            cls.svg_id = mat(s1, "x.svg", "u2.svg", b"<svg xmlns='http://www.w3.org/2000/svg'><script>1</script></svg>")
            cls.foreign_id = mat(s2, "cudzy.docx", "u3.docx", fx.docx(["TAJNE"]))

            def link(title, url):
                m = StudyMaterial(subject_id=s1.id, uploaded_by=a.id, title=title,
                                  filevault_link=url, is_filevault_link=True)
                db.session.add(m); db.session.flush()
                return m.id
            cls.fv_docx_id = link("fv-docx", "https://koloseum.example/filevault/s/DOCX")
            cls.fv_png_id = link("fv-png", "https://koloseum.example/filevault/s/PNG")
            cls.fv_folder_id = link("fv-folder", "https://koloseum.example/filevault/sf/FOLD")
            db.session.commit()
            cls.ids = dict(a=a.id, o=o.id, s1=s1.id, s2=s2.id)

    @classmethod
    def tearDownClass(cls):
        cls.fv.stop()
        cls.tmp.cleanup()

    def test_filevault_link_docx_rendered_and_png_falls_back(self):
        c = self.client()
        r = c.get(f"/subjects/{self.ids['s1']}/materials/{self.fv_docx_id}/preview-doc")
        self.assertEqual(r.status_code, 200)
        self.assertIn("Treść z FileVault", r.get_data(as_text=True))
        # obraz: odrzucony po nagłówkach (415) -> front pokaże zwykły podgląd FileVault
        r = c.get(f"/subjects/{self.ids['s1']}/materials/{self.fv_png_id}/preview-doc")
        self.assertEqual(r.status_code, 415)
        r = c.get(f"/subjects/{self.ids['s1']}/materials/{self.fv_folder_id}/preview-doc")
        self.assertEqual(r.status_code, 415)

    def test_subject_page_wires_filevault_doc_preview(self):
        html = self.client().get(f"/subjects/{self.ids['s1']}").get_data(as_text=True)
        self.assertIn(f"/materials/{self.fv_docx_id}/preview-doc", html)

    def client(self, who="a"):
        c = self.app.test_client()
        with c.session_transaction() as s:
            s["_user_id"] = str(self.ids[who]); s["_fresh"] = True
        return c

    def test_preview_doc_html_is_sandboxed(self):
        r = self.client().get(f"/subjects/{self.ids['s1']}/materials/{self.docx_id}/preview-doc")
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/html", r.content_type)
        self.assertIn("Treść wykładu", r.get_data(as_text=True))
        self.assertIn("sandbox", r.headers["Content-Security-Policy"])

    def test_preview_doc_unsupported_returns_415_page(self):
        r = self.client().get(f"/subjects/{self.ids['s1']}/materials/{self.svg_id}/preview-doc")
        self.assertEqual(r.status_code, 415)

    def test_no_cross_subject_access(self):
        c = self.client("a")
        # materiał z cudzego pokoju przez własny przedmiot — 404, nie wyciek
        for suffix in ("preview-doc", "preview", "download"):
            r = c.get(f"/subjects/{self.ids['s1']}/materials/{self.foreign_id}/{suffix}")
            self.assertEqual(r.status_code, 404, suffix)
        # cudzy przedmiot — 403
        r = c.get(f"/subjects/{self.ids['s2']}/materials/{self.foreign_id}/preview-doc")
        self.assertEqual(r.status_code, 403)

    def test_svg_served_with_sandbox_csp(self):
        r = self.client().get(f"/subjects/{self.ids['s1']}/materials/{self.svg_id}/preview")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content_type.startswith("image/svg+xml"))
        self.assertIn("sandbox", r.headers["Content-Security-Policy"])

    def test_mimetype_table_is_complete(self):
        from controllers.subject_controller import PREVIEWABLE_MIMETYPES as M
        self.assertTrue(all(M.values()), [k for k, v in M.items() if not v])
        self.assertEqual(M[".png"], "image/png"); self.assertEqual(M[".pdf"], "application/pdf")
        self.assertEqual(M[".py"], "text/plain"); self.assertEqual(M[".flac"], "audio/flac")

    def test_text_served_inline_not_as_attachment(self):
        import os
        with open(f"{self.tmp.name}/uploads/materials/u9.py", "wb") as f:
            f.write(b"print(1)")
        with self.app.app_context():
            m = StudyMaterial(subject_id=self.ids["s1"], uploaded_by=self.ids["a"], title="k.py",
                              filename="u9.py", original_filename="k.py")
            db.session.add(m); db.session.commit(); mid = m.id
        r = self.client().get(f"/subjects/{self.ids['s1']}/materials/{mid}/preview")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.headers["Content-Disposition"].startswith("inline"))
        self.assertTrue(r.content_type.startswith("text/plain"))

    def test_subject_page_renders_with_preview_buttons(self):
        r = self.client().get(f"/subjects/{self.ids['s1']}")
        self.assertEqual(r.status_code, 200)
        html = r.get_data(as_text=True)
        self.assertIn("preview-doc", html)
        self.assertIn("openPreview(", html)


if __name__ == "__main__":
    unittest.main()
